from __future__ import annotations

from datetime import date, datetime
import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from ocr_client import (
    ResponseExtractPaidEduContractDateFromImage,
    extract_paid_edu_contract_date_from_image_with_qwen3_vl,
)
from prompt_settings import ERROR_TOKEN
from transfer_excel_data import transfer_excel_data

DIRECTORY_WITH_IMAGES = Path("agreements")
FILE_PATH = Path("СВОД_СБЕР_выгрузка.xlsx")
DATA_START_ROW = 2
LOGGER = logging.getLogger(__name__)
MAX_FILES_CHECKS = 10000


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def _is_missing_or_error(value: object) -> bool:
    text = _as_text(value)
    return not text or text.upper() == str(ERROR_TOKEN).upper()


def _normalize_date_for_image_name(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")

    raw_text = _as_text(value).replace("\xa0", " ")
    if not raw_text:
        return ""

    date_candidates = [raw_text]
    first_token = raw_text.split()[0]
    if first_token != raw_text:
        date_candidates.append(first_token)

    date_formats = (
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%Y.%m.%d",
    )
    for candidate in date_candidates:
        for date_format in date_formats:
            try:
                return datetime.strptime(candidate, date_format).strftime("%d.%m.%Y")
            except ValueError:
                continue
    return raw_text


def _find_image_path(agreement_date: str, agreement_num: str) -> Path | None:
    ext_candidates = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff")
    base_names = (
        f"{agreement_date} {agreement_num}",
        f"{agreement_date}_{agreement_num}",
    )

    for base_name in base_names:
        for ext in ext_candidates:
            candidate = DIRECTORY_WITH_IMAGES / f"{base_name}{ext}"
            if candidate.exists():
                return candidate

    for pattern in (f"{agreement_date} {agreement_num}.*", f"{agreement_date}_{agreement_num}.*"):
        matches = sorted(path for path in DIRECTORY_WITH_IMAGES.glob(pattern) if path.is_file())
        if matches:
            return matches[0]
    return None


def check_paid_edu_contract_date_from_image_with_qwen3_vl() -> bool:
    """Дозаполняет 27-ю колонку (дата договора) через Qwen3-VL, если она равна ОШИБКА."""
    if not DIRECTORY_WITH_IMAGES.exists():
        LOGGER.error("Directory with images not found: %s", DIRECTORY_WITH_IMAGES)
        return False

    if not FILE_PATH.exists() and transfer_excel_data() != 0:
        LOGGER.error("Output file not found and rebuild failed: %s", FILE_PATH)
        return False

    try:
        wb_out = load_workbook(FILE_PATH)
    except TypeError as exc:
        if "extLst" not in str(exc):
            raise
        if transfer_excel_data() != 0:
            LOGGER.error("Failed to rebuild output: %s", FILE_PATH)
            return False
        wb_out = load_workbook(FILE_PATH)

    ws_out = wb_out.active
    alignment = Alignment(horizontal="center", vertical="center")
    font = Font(name="Times New Roman", size=12)

    checked_rows = 0
    updated_rows = 0
    missing_images = 0
    failed_ocr = 0
    ocr_calls = 0

    try:
        for row in ws_out.iter_rows(min_row=DATA_START_ROW):
            agreement_num_cell = row[3]  # 4-й столбец
            agreement_date_cell = row[4]  # 5-й столбец

            agreement_num = _as_text(agreement_num_cell.value)
            agreement_date = _normalize_date_for_image_name(agreement_date_cell.value)
            if not agreement_num or not agreement_date:
                continue

            university_name = row[22].value  # 23
            student_fio = row[23].value  # 24
            customer_fio = row[24].value  # 25
            paid_edu_contract_number = row[25].value  # 26
            paid_edu_contract_date = row[26].value  # 27
            specialty_code = row[27].value  # 28

            has_any_extracted_context = any(
                not _is_missing_or_error(value)
                for value in (
                    university_name,
                    student_fio,
                    customer_fio,
                    paid_edu_contract_number,
                    specialty_code,
                )
            )
            needs_date_recheck = _is_missing_or_error(paid_edu_contract_date)
            if not (has_any_extracted_context and needs_date_recheck):
                continue

            checked_rows += 1
            image_path = _find_image_path(agreement_date, agreement_num)
            if image_path is None:
                missing_images += 1
                LOGGER.warning(
                    "Image not found for agreement: date=%s number=%s",
                    agreement_date,
                    agreement_num,
                )
                continue

            if ocr_calls >= MAX_FILES_CHECKS:
                LOGGER.info("Reached OCR limit: %d", MAX_FILES_CHECKS)
                break

            response: ResponseExtractPaidEduContractDateFromImage = (
                extract_paid_edu_contract_date_from_image_with_qwen3_vl(str(image_path))
            )
            ocr_calls += 1
            if response.success and response.date:
                target_cell = ws_out.cell(row=agreement_num_cell.row, column=27, value=response.date)
                target_cell.alignment = alignment
                target_cell.font = font
                updated_rows += 1
            else:
                failed_ocr += 1

        wb_out.save(FILE_PATH)
    finally:
        wb_out.close()

    LOGGER.info(
        "Date recheck completed: checked=%d, updated=%d, missing_images=%d, failed_ocr=%d, ocr_calls=%d",
        checked_rows,
        updated_rows,
        missing_images,
        failed_ocr,
        ocr_calls,
    )
    return True


if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    raise SystemExit(0 if check_paid_edu_contract_date_from_image_with_qwen3_vl() else 1)
