from __future__ import annotations

from datetime import date, datetime
import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from prompt_settings import ERROR_TOKEN

LOGGER = logging.getLogger(__name__)

FILE_PATH = Path("СВОД_СБЕР_выгрузка.xlsx")
DATA_START_ROW = 3
MAX_MISMATCH_WARNINGS = 10

REPORTING_DATES_BY_MONTH = {
    "01": "ЯНВАРЬ!",
    "02": "ФЕВРАЛЬ!",
    "03": "МАРТ!",
    "04": "АПРЕЛЬ!",
    "05": "МАЙ!",
    "06": "ИЮНЬ!",
    "07": "ИЮЛЬ;АВГУСТ;СЕНТЯБРЬ",
    "08": "АВГУСТ;СЕНТЯБРЬ",
    "09": "СЕНТЯБРЬ",
    "10": "ОКТЯБРЬ!",
    "11": "НОЯБРЬ!",
    "12": "ДЕКАБРЬ!",
}


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def _is_missing_or_error(value: object) -> bool:
    text = _as_text(value)
    return not text or text.upper() == str(ERROR_TOKEN).upper()


def _extract_month(value: object) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%m")

    text = _as_text(value).replace("\xa0", " ")
    if not text or text.upper() == str(ERROR_TOKEN).upper():
        return None

    if text.isdigit() and 1 <= int(text) <= 12:
        return f"{int(text):02d}"

    candidates = [text]
    first_token = text.split()[0]
    if first_token != text:
        candidates.append(first_token)

    date_formats = (
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%Y.%m.%d",
    )
    for candidate in candidates:
        for date_format in date_formats:
            try:
                parsed = datetime.strptime(candidate, date_format)
                return parsed.strftime("%m")
            except ValueError:
                continue

    return None


def check_data_9col(overwrite_existing: bool = False) -> bool:
    """
    Проверяет 27-ю колонку (дата договора) и выставляет в 9-й колонке месяц/месяцы по правилу.

    По умолчанию заполняет только пустые ячейки 9-й колонки.
    """
    if not FILE_PATH.exists():
        LOGGER.error("File not found: %s", FILE_PATH)
        return False

    try:
        wb_out = load_workbook(FILE_PATH)
    except TypeError as exc:
        if "extLst" not in str(exc):
            raise
        wb_out = load_workbook(FILE_PATH)

    ws_out = wb_out.active
    alignment = Alignment(horizontal="center", vertical="center")
    font = Font(name="Times New Roman", size=12)

    rows_total = 0
    updated_rows = 0
    skipped_existing = 0
    missing_or_error_dates = 0
    invalid_date_values = 0
    invalid_month_values = 0
    mismatched_existing = 0

    try:
        for row in ws_out.iter_rows(min_row=DATA_START_ROW):
            rows_total += 1
            col9_cell = row[8]
            col27_cell = row[26]

            date_value = col27_cell.value
            if _is_missing_or_error(date_value):
                missing_or_error_dates += 1
                continue

            month = _extract_month(date_value)
            if month is None:
                invalid_date_values += 1
                LOGGER.warning(
                    "Cannot parse date in row=%s, col27=%r",
                    col27_cell.row,
                    date_value,
                )
                continue

            reporting_dates = REPORTING_DATES_BY_MONTH.get(month)
            if reporting_dates is None:
                invalid_month_values += 1
                LOGGER.warning(
                    "Month is out of mapping in row=%s, col27=%r, month=%s",
                    col27_cell.row,
                    date_value,
                    month,
                )
                continue

            current_col9_value = _as_text(col9_cell.value)
            if current_col9_value and current_col9_value != reporting_dates:
                mismatched_existing += 1
                if mismatched_existing <= MAX_MISMATCH_WARNINGS:
                    LOGGER.warning(
                        "Col9 mismatch in row=%s: current=%r, expected=%r from col27=%r",
                        col9_cell.row,
                        current_col9_value,
                        reporting_dates,
                        date_value,
                    )

            if current_col9_value and not overwrite_existing:
                skipped_existing += 1
                continue

            if current_col9_value != reporting_dates:
                col9_cell.value = reporting_dates
                col9_cell.alignment = alignment
                col9_cell.font = font
                updated_rows += 1

        wb_out.save(FILE_PATH)
    finally:
        wb_out.close()

    if mismatched_existing > MAX_MISMATCH_WARNINGS:
        LOGGER.warning(
            "Suppressed %d additional col9 mismatch warnings.",
            mismatched_existing - MAX_MISMATCH_WARNINGS,
        )

    LOGGER.info(
        (
            "9th column update completed: rows_total=%d, updated=%d, skipped_existing=%d, "
            "missing_or_error_dates=%d, invalid_date_values=%d, invalid_month_values=%d, "
            "mismatched_existing=%d"
        ),
        rows_total,
        updated_rows,
        skipped_existing,
        missing_or_error_dates,
        invalid_date_values,
        invalid_month_values,
        mismatched_existing,
    )
    return True


if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    raise SystemExit(0 if check_data_9col() else 1)
