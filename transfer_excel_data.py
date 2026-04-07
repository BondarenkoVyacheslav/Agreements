from __future__ import annotations

from copy import copy
from datetime import date, datetime
import logging
from pathlib import Path
import re
from shutil import copy2

import dspy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from extract_fields import ExtractedFields
from ocr_client import check_university_name_llm
from prompt_settings import ERROR_TOKEN
from utils import _norm as _norm_text, _norm_code

LOGGER = logging.getLogger(__name__)

TEMPLATE_PATH = Path("1448 декабрь в Сбер на проверку Шаблон.xlsx")
HEADER_STYLE_TEMPLATE_PATH = Path("Шаблон_выгрузки.xlsx")
OUTPUT_PATH = Path("1448 декабрь в Сбер на проверку Выгрузка.xlsx")
WORKSHEET_NAME = "CustomQuery"
DATA_START_ROW = 2

EXCEL_CONTRACT_COLUMN = 5  # E
BASE_UNIVERSITY_COLUMN = 3  # C
BASE_STUDENT_COLUMN = 4  # D
BASE_SPECIALTY_COLUMN = 8  # H

AI_OUTPUT_START_COLUMN = 11  # K
AI_HEADERS_SOURCE_START_COLUMN = 23  # W
AI_HEADERS_COUNT = 7
AI_RESULT_COLUMN = AI_OUTPUT_START_COLUMN + AI_HEADERS_COUNT - 1  # Q

NUMBER_FIRST_STEM_PATTERN = re.compile(r"^(?P<number>.+?)[_ ]+(?P<date>\d{1,2}\.\d{1,2}\.\d{4})$")
DATE_FIRST_STEM_PATTERN = re.compile(r"^(?P<date>\d{1,2}\.\d{1,2}\.\d{4})[_ ]+(?P<number>.+?)$")
DATE_ONLY_PATTERN = re.compile(r"^\d{1,2}\.\d{1,2}\.\d{4}$")
DASH_PATTERN = re.compile(r"[‐‑‒–—―−]")


def _get_output_worksheet(workbook):
    if WORKSHEET_NAME in workbook.sheetnames:
        return workbook[WORKSHEET_NAME]
    return workbook.active


def _copy_cell_style(source_cell, target_cell) -> None:
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def _safe_value(value: object) -> str:
    if value is None:
        return ERROR_TOKEN
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text if text else ERROR_TOKEN


def _load_output_workbook():
    try:
        return load_workbook(OUTPUT_PATH)
    except TypeError as exc:
        if "extLst" not in str(exc):
            raise
        if transfer_excel_data() != 0:
            raise RuntimeError(f"Failed to rebuild output workbook: {OUTPUT_PATH}") from exc
        return load_workbook(OUTPUT_PATH)


def _prepare_output_headers(ws_out, ws_style) -> None:
    body_style_source = ws_out.cell(row=DATA_START_ROW, column=AI_OUTPUT_START_COLUMN)

    for offset in range(AI_HEADERS_COUNT):
        source_column = AI_HEADERS_SOURCE_START_COLUMN + offset
        target_column = AI_OUTPUT_START_COLUMN + offset

        source_header_cell = ws_style.cell(row=1, column=source_column)
        target_header_cell = ws_out.cell(row=1, column=target_column, value=source_header_cell.value)
        _copy_cell_style(source_header_cell, target_header_cell)

        source_letter = get_column_letter(source_column)
        target_letter = get_column_letter(target_column)
        ws_out.column_dimensions[target_letter].width = ws_style.column_dimensions[source_letter].width

        for row_index in range(DATA_START_ROW, ws_out.max_row + 1):
            body_cell = ws_out.cell(row=row_index, column=target_column)
            _copy_cell_style(body_style_source, body_cell)

    ws_out.row_dimensions[1].height = ws_style.row_dimensions[1].height
    ws_out.auto_filter.ref = f"A1:{get_column_letter(AI_RESULT_COLUMN)}1"


def normalize_credit_contract_number(value: object) -> str:
    if value is None:
        return ""

    normalized = str(value).strip().upper().rstrip(".")
    if not normalized:
        return ""

    normalized = DASH_PATTERN.sub("-", normalized)
    normalized = normalized.replace("/", "-").replace("\\", "-").replace("_", "-")
    normalized = re.sub(r"\s+", "-", normalized)
    normalized = re.sub(r"-{2,}", "-", normalized)
    return normalized.strip("-.")


def extract_credit_contract_number_from_excel_cell(value: object) -> str | None:
    if value is None:
        return None

    parts = str(value).split()
    if len(parts) < 2:
        return None

    contract_number = normalize_credit_contract_number(parts[1])
    return contract_number or None


def extract_credit_contract_number_from_document_name(path: Path | str) -> str | None:
    stem = Path(path).stem.rstrip(".").strip()
    if not stem or DATE_ONLY_PATTERN.fullmatch(stem) or not re.search(r"\d", stem):
        return None

    for pattern in (NUMBER_FIRST_STEM_PATTERN, DATE_FIRST_STEM_PATTERN):
        match = pattern.fullmatch(stem)
        if match:
            contract_number = normalize_credit_contract_number(match.group("number"))
            return contract_number or None

    contract_number = normalize_credit_contract_number(stem)
    return contract_number or None


def find_matching_rows(ws, contract_number: str) -> list[int]:
    normalized_contract_number = normalize_credit_contract_number(contract_number)
    if not normalized_contract_number:
        return []

    matching_rows: list[int] = []
    for row_index in range(DATA_START_ROW, ws.max_row + 1):
        cell_value = ws.cell(row=row_index, column=EXCEL_CONTRACT_COLUMN).value
        row_contract_number = extract_credit_contract_number_from_excel_cell(cell_value)
        if row_contract_number == normalized_contract_number:
            matching_rows.append(row_index)

    return matching_rows


def transfer_excel_data() -> int:
    if not TEMPLATE_PATH.exists():
        print(f"File not found: {TEMPLATE_PATH}")
        return 1
    if not HEADER_STYLE_TEMPLATE_PATH.exists():
        print(f"File not found: {HEADER_STYLE_TEMPLATE_PATH}")
        return 1

    copy2(TEMPLATE_PATH, OUTPUT_PATH)

    wb_out = load_workbook(OUTPUT_PATH)
    wb_style = load_workbook(HEADER_STYLE_TEMPLATE_PATH)
    try:
        ws_out = _get_output_worksheet(wb_out)
        ws_style = wb_style.active
        _prepare_output_headers(ws_out, ws_style)
        wb_out.save(OUTPUT_PATH)
    finally:
        wb_style.close()
        wb_out.close()

    return 0


def transfer_reporting_data_to_excel(*_args, **_kwargs) -> bool:
    LOGGER.warning("Логика reporting_dates отключена для новой декабрьской выгрузки.")
    return False


def _check_university_match(cell_name: str, ex_name: str, lm: dspy.LM | None = None) -> tuple[bool, str]:
    if not cell_name or not ex_name:
        return False, "Пустое название"

    if cell_name.upper() == ERROR_TOKEN or ex_name.upper() == ERROR_TOKEN:
        return False, "ОШИБКА в названии"

    cell_norm = _norm_text(cell_name)
    ex_norm = _norm_text(ex_name)

    def _normalize_university_for_compare(value: str) -> str:
        normalized = _norm_text(value)
        normalized = re.sub(r"[^0-9A-ZА-Я]+", " ", normalized)
        normalized = re.sub(r"\s+", " ", normalized).strip()
        return normalized

    def _quoted_core(value: str) -> str:
        match = re.search(r"[«\"]([^»\"]+)[»\"]", str(value))
        if not match:
            return ""
        return _norm_text(match.group(1))

    if cell_norm == ex_norm:
        return True, "Точное совпадение"

    if cell_norm in ex_norm or ex_norm in cell_norm:
        return True, "Одно в другом"

    cell_soft_norm = _normalize_university_for_compare(cell_name)
    ex_soft_norm = _normalize_university_for_compare(ex_name)
    if cell_soft_norm == ex_soft_norm:
        return True, "Совпадение после нормализации"
    if cell_soft_norm in ex_soft_norm or ex_soft_norm in cell_soft_norm:
        return True, "Одно в другом после нормализации"

    cell_quoted_core = _quoted_core(cell_name)
    ex_quoted_core = _quoted_core(ex_name)
    if cell_quoted_core and ex_quoted_core and cell_quoted_core == ex_quoted_core:
        return True, "Совпадение по названию в кавычках"

    try:
        llm_result = check_university_name_llm(cell_name, ex_name, lm)
        LOGGER.info("Уточняем у LLM сходство названий вузов.")
        if llm_result:
            return True, "Совпадение по LLM"
        return False, "Разные ВУЗы (LLM)"
    except Exception as exc:
        LOGGER.warning("Ошибка LLM при сравнении ВУЗов: %s", exc)
        return False, "Ошибка LLM"


def _build_ai_result(ws_out, row_index: int, ex_university: str, ex_student_fio: str, ex_specialty: str, lm: dspy.LM | None) -> str:
    cell_university_name = ws_out.cell(row=row_index, column=BASE_UNIVERSITY_COLUMN).value
    cell_fio = ws_out.cell(row=row_index, column=BASE_STUDENT_COLUMN).value
    cell_cnp = ws_out.cell(row=row_index, column=BASE_SPECIALTY_COLUMN).value

    messages: list[str] = []

    cell_university_norm = _norm_text(str(cell_university_name) if cell_university_name is not None else "")
    ex_university_norm = _norm_text(ex_university)
    if cell_university_norm not in ("", ERROR_TOKEN) and ex_university_norm not in ("", ERROR_TOKEN):
        is_match, _comment = _check_university_match(str(cell_university_name), ex_university, lm)
        if not is_match:
            messages.append("Иной вуз")

    cell_fio_norm = _norm_text(str(cell_fio) if cell_fio is not None else "")
    ex_fio_norm = _norm_text(ex_student_fio)
    if cell_fio_norm not in ("", ERROR_TOKEN) and ex_fio_norm not in ("", ERROR_TOKEN):
        cell_parts = cell_fio_norm.split()
        ex_parts = ex_fio_norm.split()

        cell_surname = cell_parts[0] if len(cell_parts) >= 1 else ""
        cell_name = cell_parts[1] if len(cell_parts) >= 2 else ""
        cell_patronymic = cell_parts[2] if len(cell_parts) >= 3 else ""

        ex_surname = ex_parts[0] if len(ex_parts) >= 1 else ""
        ex_name = ex_parts[1] if len(ex_parts) >= 2 else ""
        ex_patronymic = ex_parts[2] if len(ex_parts) >= 3 else ""

        surname_match = cell_surname == ex_surname
        name_match = cell_name == ex_name
        patronymic_match = cell_patronymic == ex_patronymic
        matches_count = sum([surname_match, name_match, patronymic_match])

        fio_differences: list[str] = []
        if not surname_match:
            fio_differences.append("Иная фамилия")
        if not name_match:
            fio_differences.append("Иное имя")
        if not patronymic_match and cell_patronymic and ex_patronymic:
            pass

        if matches_count >= 2:
            if fio_differences:
                messages.append("; ".join(fio_differences))
        elif matches_count == 1:
            messages.append("Не найден")
            if fio_differences:
                messages.append("; ".join(fio_differences))
        else:
            messages.append("Не найден")

    cell_cnp_norm = _norm_code(str(cell_cnp) if cell_cnp is not None else "")
    ex_specialty_norm = _norm_code(ex_specialty)
    if cell_cnp_norm not in ("", ERROR_TOKEN) and ex_specialty_norm not in ("", ERROR_TOKEN):
        if cell_cnp_norm != ex_specialty_norm:
            messages.append("Иное НПС")

    if "Иное НПС" in messages:
        messages.append(f"={ex_specialty}")

    return "; ".join(messages)


def transfer_extracted_data_and_logic_to_excel(
    edu_loan_agr_num: str,
    exctracted_fields: ExtractedFields,
    lm: dspy.LM | None = None,
) -> bool:
    if not OUTPUT_PATH.exists():
        if transfer_excel_data() != 0:
            print(f"File not found: {OUTPUT_PATH}")
            return False

    normalized_contract_number = normalize_credit_contract_number(edu_loan_agr_num)
    if not normalized_contract_number:
        return False

    wb_out = _load_output_workbook()
    try:
        ws_out = _get_output_worksheet(wb_out)
        matching_rows = find_matching_rows(ws_out, normalized_contract_number)
        if not matching_rows:
            return False

        body_style_source = ws_out.cell(row=DATA_START_ROW, column=AI_OUTPUT_START_COLUMN)

        ex_university = _safe_value(exctracted_fields.university_name)
        ex_student_fio = _safe_value(exctracted_fields.student_fio)
        ex_customer_fio = _safe_value(exctracted_fields.customer_fio)
        ex_paid_number = _safe_value(exctracted_fields.paid_edu_contract_number)
        ex_paid_date = _safe_value(exctracted_fields.paid_edu_contract_date)
        ex_specialty = _safe_value(exctracted_fields.specialty_code)

        extracted_values = [
            ex_university,
            ex_student_fio,
            ex_customer_fio,
            ex_paid_number,
            ex_paid_date,
            ex_specialty,
        ]

        for row_index in matching_rows:
            for offset, value in enumerate(extracted_values):
                cell = ws_out.cell(row=row_index, column=AI_OUTPUT_START_COLUMN + offset, value=value)
                _copy_cell_style(body_style_source, cell)

            ai_result = _build_ai_result(ws_out, row_index, ex_university, ex_student_fio, ex_specialty, lm)
            result_cell = ws_out.cell(row=row_index, column=AI_RESULT_COLUMN, value=ai_result)
            _copy_cell_style(body_style_source, result_cell)

        wb_out.save(OUTPUT_PATH)
        return True
    finally:
        wb_out.close()


if __name__ == "__main__":
    raise SystemExit(transfer_excel_data())
